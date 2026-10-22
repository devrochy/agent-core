"""Ingesta del pack de conocimiento sin base (ADR-0145 §2, HU-AOS-49 parte 2).

- Tipo por MIME o extensión, confirmado por el contenido (magic bytes, ZIP de Word, texto).
- Rutas del original dentro de la carpeta del tenant, sin traversal.
- Extracción de PDF, DOCX, TXT y Markdown con sección y página; un PDF sin texto es `sin_texto`.
- Troceado v1: no mezcla secciones, respeta el objetivo y el máximo, solapa y no parte frases.
"""

from __future__ import annotations

import io
import uuid
import zipfile

import pytest
from _pdf_minimo import pdf_minimo as _pdf
from pack_knowledge import archivos
from pack_knowledge.archivos import MIME_DOCX, MIME_MD, MIME_PDF, MIME_TXT, ArchivoInvalido
from pack_knowledge.extraccion import Bloque, ExtraccionError, extraer_texto
from pack_knowledge.troceado import (
    MAX_FRAGMENTOS,
    MAX_TOKENS,
    OBJETIVO_TOKENS,
    DemasiadosFragmentos,
    tokens_estimados,
    trocear,
)


def _docx() -> bytes:
    from docx import Document

    documento = Document()
    documento.add_heading("Envíos", level=1)
    documento.add_paragraph("Hacemos envíos a todo el país en 3 días hábiles.")
    documento.add_heading("Cambios", level=1)
    documento.add_paragraph("Aceptamos cambios hasta 30 días después de la compra.")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "Camiseta"
    tabla.rows[0].cells[1].text = "$45.000"
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


# ------------------------------------------------------------------ tipo y contenido
@pytest.mark.parametrize(
    ("mime", "nombre", "esperado"),
    [
        ("application/pdf", "faq.pdf", MIME_PDF),
        ("application/octet-stream", "notas.md", MIME_MD),
        ("text/x-markdown", "notas.md", MIME_MD),
        ("text/plain; charset=utf-8", "faq.txt", MIME_TXT),
        ("", "politica.docx", MIME_DOCX),
    ],
)
def test_tipo_por_mime_o_extension(mime: str, nombre: str, esperado: str) -> None:
    assert archivos.tipo_de(mime, nombre) == esperado


@pytest.mark.parametrize(("mime", "nombre"), [("image/png", "foto.png"), ("application/zip", "x.zip"), ("", "x.exe")])
def test_tipo_no_admitido(mime: str, nombre: str) -> None:
    with pytest.raises(ArchivoInvalido) as exc:
        archivos.tipo_de(mime, nombre)
    assert exc.value.codigo == "formato"


def test_contenido_confirma_el_tipo() -> None:
    archivos.validar(MIME_PDF, _pdf(["hola"]))
    archivos.validar(MIME_DOCX, _docx())
    archivos.validar(MIME_TXT, "Envíos gratis".encode())
    archivos.validar(MIME_TXT, "Envíos gratis".encode("latin-1"))
    zip_cualquiera = io.BytesIO()
    with zipfile.ZipFile(zip_cualquiera, "w") as z:
        z.writestr("hola.txt", "x")
    for tipo, contenido in (
        (MIME_PDF, b"no soy un pdf"),
        (MIME_DOCX, zip_cualquiera.getvalue()),
        (MIME_TXT, b"\x89PNG\r\n\x1a\n\x00\x00binario"),
    ):
        with pytest.raises(ArchivoInvalido) as exc:
            archivos.validar(tipo, contenido)
        assert exc.value.codigo == "formato", tipo


def test_vacio_y_muy_grande() -> None:
    with pytest.raises(ArchivoInvalido) as exc:
        archivos.validar(MIME_TXT, b"")
    assert exc.value.codigo == "vacio"
    with pytest.raises(ArchivoInvalido) as exc:
        archivos.validar(MIME_TXT, b"a" * (archivos.MAX_BYTES + 1))
    assert exc.value.codigo == "muy_grande"


# ------------------------------------------------------------------ rutas
def test_guardar_resolver_y_borrar(tmp_path, monkeypatch) -> None:
    import core.almacenamiento as almacenamiento

    monkeypatch.setattr(archivos, "carpeta_del_tenant", lambda slug: tmp_path / slug)
    # El alta pasa por el controlador único: se aísla también su raíz.
    monkeypatch.setattr(almacenamiento, "carpeta_del_tenant", lambda slug: tmp_path / slug)
    monkeypatch.setattr(almacenamiento, "tenants_root", lambda: tmp_path)
    doc = uuid.uuid4()
    ruta = archivos.guardar("negocio-a", doc, MIME_PDF, b"%PDF-1.4 x")
    assert ruta == f"{doc}.pdf"
    assert archivos.resolver("negocio-a", ruta) == (tmp_path / "negocio-a" / "conocimiento" / ruta).resolve()
    # Otro negocio, traversal y rutas absolutas no resuelven.
    (tmp_path / "negocio-b" / "conocimiento").mkdir(parents=True)
    (tmp_path / "negocio-b" / "conocimiento" / "x.pdf").write_bytes(b"%PDF")
    for intento in ("../../negocio-b/conocimiento/x.pdf", "/etc/passwd", f"sub/{ruta}", "", None):
        assert archivos.resolver("negocio-a", intento) is None, intento
    archivos.borrar("negocio-a", ruta)
    assert archivos.resolver("negocio-a", ruta) is None


# ------------------------------------------------------------------ extracción
def test_pdf_con_paginas() -> None:
    extraido = extraer_texto(MIME_PDF, _pdf(["Hacemos envios a todo el pais", "Cambios hasta 30 dias"]))
    assert extraido.paginas == 2
    assert [(b.pagina, b.texto) for b in extraido.bloques] == [
        (1, "Hacemos envios a todo el pais"),
        (2, "Cambios hasta 30 dias"),
    ]


def test_pdf_sin_texto_y_roto() -> None:
    with pytest.raises(ExtraccionError) as exc:
        extraer_texto(MIME_PDF, _pdf([""]))
    assert exc.value.codigo == "sin_texto"
    with pytest.raises(ExtraccionError) as exc:
        extraer_texto(MIME_PDF, b"%PDF-1.4 roto")
    assert exc.value.codigo == "formato"


def test_docx_con_secciones_y_tablas() -> None:
    bloques = extraer_texto(MIME_DOCX, _docx()).bloques
    assert Bloque("Hacemos envíos a todo el país en 3 días hábiles.", "Envíos") in bloques
    assert Bloque("Aceptamos cambios hasta 30 días después de la compra.", "Cambios") in bloques
    assert Bloque("Camiseta | $45.000", "Tablas") in bloques


def test_markdown_con_encabezados_y_txt_plano() -> None:
    md = b"# Horario\nLunes a viernes\nde 8 a 5.\n\n## Domingos\nCerrado, salvo festivos.\n"
    assert extraer_texto(MIME_MD, md).bloques == [
        Bloque("Lunes a viernes de 8 a 5.", "Horario"),
        Bloque("Cerrado, salvo festivos.", "Domingos"),
    ]
    # En TXT un «#» es texto.
    assert extraer_texto(MIME_TXT, b"# no es titulo\nsigue el parrafo").bloques == [
        Bloque("# no es titulo sigue el parrafo")
    ]


def test_texto_casi_vacio_es_sin_texto() -> None:
    with pytest.raises(ExtraccionError) as exc:
        extraer_texto(MIME_TXT, b"  hola  \n\n")
    assert exc.value.codigo == "sin_texto"


# ------------------------------------------------------------------ troceado
def _frase(i: int) -> str:
    return f"La frase número {i} habla de envíos, cambios y horarios del negocio."


def test_troceado_respeta_tamanos_y_solapa() -> None:
    texto = " ".join(_frase(i) for i in range(400))
    fragmentos = trocear([Bloque(texto, "FAQ", 1)])
    assert len(fragmentos) > 3
    assert [f.ordinal for f in fragmentos] == list(range(len(fragmentos)))
    for f in fragmentos:
        assert f.tokens <= MAX_TOKENS
        assert f.seccion == "FAQ"
        assert f.contenido.endswith(".")  # sin partir frases
    assert all(f.tokens >= OBJETIVO_TOKENS * 0.8 for f in fragmentos[:-1])
    # Solape: el siguiente arranca con las últimas frases enteras del anterior (≤ 60 tokens).
    primera_del_siguiente = fragmentos[1].contenido.split(". ", 1)[0] + "."
    assert primera_del_siguiente in fragmentos[0].contenido
    cola = fragmentos[0].contenido[fragmentos[0].contenido.index(primera_del_siguiente) :]
    assert fragmentos[1].contenido.startswith(cola)
    assert tokens_estimados(cola) <= 60


def test_troceado_no_mezcla_secciones_ni_solapa_entre_ellas() -> None:
    fragmentos = trocear([Bloque("Envíos gratis desde 100.000.", "Envíos", 1), Bloque("Cambios en 30 días.", "Cambios", 2)])
    assert [(f.seccion, f.contenido, f.pagina_desde, f.pagina_hasta) for f in fragmentos] == [
        ("Envíos", "Envíos gratis desde 100.000.", 1, 1),
        ("Cambios", "Cambios en 30 días.", 2, 2),
    ]


def test_troceado_parrafo_enorme_sin_puntos() -> None:
    palabra_larga = " ".join(["palabra"] * 2_000)
    fragmentos = trocear([Bloque(palabra_larga)])
    assert all(f.tokens <= MAX_TOKENS for f in fragmentos)
    assert sum(f.contenido.count("palabra") for f in fragmentos) >= 2_000


def test_troceado_tope_de_fragmentos() -> None:
    bloques = [Bloque(_frase(i), f"S{i}") for i in range(MAX_FRAGMENTOS + 1)]
    with pytest.raises(DemasiadosFragmentos):
        trocear(bloques)


def test_tokens_estimados() -> None:
    assert tokens_estimados("abcd" * 10) == 10
    assert tokens_estimados("") == 1


# ------------------------------------------------------------------ core: PII y tareas
def test_contar_pii_no_enmascara() -> None:
    from core.memory.vector import contar_pii

    texto = "Escríbenos a ventas@tienda.co o al +57 300 123 4567."
    assert contar_pii(texto) == 2
    assert contar_pii("Envíos gratis desde 100.000 pesos") == 0


def test_tareas_registradas_con_nombre_del_pack() -> None:
    import pack_knowledge.jobs  # noqa: F401
    from pack_knowledge import ingesta

    from core.tasks.celery_app import celery

    assert {ingesta.TAREA_EXTRAER, ingesta.TAREA_EMBEBER} <= set(celery.tasks)


async def test_tarea_sin_tenant_falla() -> None:
    from pack_knowledge import ingesta

    with pytest.raises(TypeError):
        await ingesta.tarea_extraer(document_id=str(uuid.uuid4()))  # type: ignore[call-arg]


def test_register_task_registra_solo_la_clase_del_error(monkeypatch) -> None:
    from core.plugin_manager import plugin_manager
    from core.tasks import decorators

    async def _nada() -> None:
        return None

    monkeypatch.setattr(plugin_manager, "ensure_process_activation", _nada)
    monkeypatch.setattr(plugin_manager, "is_active", lambda _pid: True)
    vistos: list[tuple[str, dict]] = []
    monkeypatch.setattr(decorators.log, "error", lambda evento, **kw: vistos.append((evento, kw)))
    monkeypatch.setattr(decorators.log, "exception", lambda *a, **k: vistos.append(("exception", k)))

    @decorators.register_task("pack_prueba_49")
    async def explota(*, texto: str) -> None:
        raise ValueError(f"parámetros: {texto}")

    decorators.celery.tasks["pack_prueba_49.explota"](texto="contenido secreto del documento")
    assert vistos == [("task_failed", {"plugin": "pack_prueba_49", "task": "pack_prueba_49.explota", "error": "ValueError"})]


# ------------------------------------------------------------------ core: carpeta del tenant
@pytest.fixture
def raiz(tmp_path, monkeypatch):
    from core import onboarding

    monkeypatch.setattr(onboarding, "tenants_root", lambda: tmp_path)
    return tmp_path


def test_borrar_carpeta_la_aparta_y_la_borra(raiz) -> None:
    from core.onboarding import borrar_carpeta_del_tenant

    (raiz / "negocio" / "playbooks").mkdir(parents=True)
    (raiz / "negocio" / "playbooks" / "ventas.md").write_text("x", encoding="utf-8")
    assert borrar_carpeta_del_tenant("negocio") is True
    assert not (raiz / "negocio").exists()
    assert list((raiz / ".purgados").iterdir()) == []
    assert borrar_carpeta_del_tenant("negocio") is False


@pytest.mark.parametrize("slug", ["../fuera", "", ".purgados", "A"])
def test_carpeta_del_tenant_rechaza_slugs_invalidos(raiz, slug: str) -> None:
    from core.onboarding import carpeta_del_tenant

    with pytest.raises(ValueError):
        carpeta_del_tenant(slug)


def test_cuarentena_solo_datos_de_clientes(raiz) -> None:
    from core.onboarding import cuarentena_de_datos_huerfanos

    for sub in ("conocimiento", "evidencias", "docs"):
        (raiz / "negocio" / sub).mkdir(parents=True)
        (raiz / "negocio" / sub / "a.txt").write_text(sub, encoding="utf-8")
    movidas = cuarentena_de_datos_huerfanos("negocio")
    assert sorted(m.name.split("-")[1] for m in movidas) == ["conocimiento", "evidencias"]
    assert (raiz / "negocio" / "docs" / "a.txt").is_file()
    assert not (raiz / "negocio" / "conocimiento").exists()
    assert cuarentena_de_datos_huerfanos("negocio") == []


async def test_purga_con_disco_fallando_alerta_y_devuelve_false(monkeypatch) -> None:
    from core import onboarding, tenants
    from core.shared import alerts

    def _falla(_slug: str) -> bool:
        raise OSError("disco")

    alertas: list[dict] = []

    async def _alerta(titulo: str, cuerpo: str, **kw) -> None:
        alertas.append({"titulo": titulo, **kw})

    errores: list[tuple[str, dict]] = []
    monkeypatch.setattr(onboarding, "borrar_carpeta_del_tenant", _falla)
    monkeypatch.setattr(alerts, "send_alert", _alerta)
    monkeypatch.setattr(tenants.log, "error", lambda evento, **kw: errores.append((evento, kw)))
    tid = uuid.uuid4()
    assert await tenants._borrar_archivos_del_tenant(tid, "negocio") is False
    assert errores == [("tenant_purge_files_cleanup_failed", {"tenant_id": str(tid), "error": "OSError"})]
    assert alertas and alertas[0]["severity"] == "critical"


# ------------------------------------------------------------------ worker: original ausente
async def test_original_ausente_falla_como_interno_no_como_formato(monkeypatch) -> None:
    from contextlib import asynccontextmanager
    from types import SimpleNamespace

    from pack_knowledge import ingesta

    doc = SimpleNamespace(estado="pendiente", mime=MIME_PDF, sha256="x", ruta_archivo="a.pdf")

    @asynccontextmanager
    async def _sesion():
        yield None

    async def _cargar(_session, _tid, _did):
        return doc

    async def _slug(_tid):
        return "negocio"

    avisos: list[tuple[str, dict]] = []
    monkeypatch.setattr(ingesta, "get_session", _sesion)
    monkeypatch.setattr(ingesta, "_cargar_para_tarea", _cargar)
    monkeypatch.setattr(ingesta, "_slug", _slug)
    monkeypatch.setattr(ingesta.archivos, "resolver", lambda _slug, _ruta: None)
    monkeypatch.setattr(ingesta.log, "warning", lambda evento, **kw: avisos.append((evento, kw)))

    await ingesta._extraer(uuid.uuid4(), uuid.uuid4())

    assert (doc.estado, doc.error_codigo, doc.error_detalle) == ("fallido", "interno", "archivo_ausente")
    assert [(e, kw["codigo"], kw["detalle"]) for e, kw in avisos] == [
        ("knowledge_documento_fallido", "interno", "archivo_ausente")
    ]
