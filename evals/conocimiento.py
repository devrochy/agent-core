"""Golden set de recuperación del pack de conocimiento (ADR-0145 §7, HU-AOS-49 parte 3).

Ingiere tres documentos de un negocio ficticio por el camino real (subida → extracción → troceado →
embeddings) y mide la búsqueda con las 50 preguntas de `evals/conocimiento/preguntas.json`:

- `recall@k`: fracción de las 40 preguntas con respuesta cuyo fragmento esperado (el que contiene
  `clave`, del documento indicado) sale entre los `k` primeros.
- `mrr`: promedio de `1/rango` de ese fragmento (0 si no sale).
- `negativas`: cuántas de las 10 preguntas que no están en los documentos devuelven cero fragmentos.

Los documentos se arman en tres formatos desde las fuentes de `evals/conocimiento/documentos/`:
- las preguntas frecuentes van en PDF, una sección por página;
- la política de cambios va en DOCX, con los encabezados como estilos de título;
- el catálogo va en TXT.

Embeddings:
- **replay** (CI y develop): `evals/conocimiento/grabacion.json`. Un texto sin grabar falla y no
  sale a la red.
- **record** (`--grabar`): llamada real a `text-embedding-3-small`, ≈ 50.000 tokens ≈ US$0,001.
  La hace el dueño o se hace con su OK (memoria del proyecto: costo de LLM).
- **hash** (`--hash`): sin grabación ni red. Solo sirve para la parte léxica.

Uso (con el stack de integración y `APP_ENV=development`):

    python evals/conocimiento.py                 # replay
    python evals/conocimiento.py --hash          # sin grabación
    OPENAI_API_KEY=... python evals/conocimiento.py --grabar
"""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import os
import re
import sys
import unicodedata
import uuid
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any

DIR = Path(__file__).resolve().parent / "conocimiento"
GRABACION = DIR / "grabacion.json"
UMBRALES = {"recall": 0.90, "mrr": 0.70, "negativas": 8}
K = 4


@dataclass(frozen=True)
class Metricas:
    modo: str
    recall: float
    mrr: float
    negativas: int
    total_negativas: int
    fallidas: list[str]
    negativas_con_resultado: list[str]

    def cumple(self) -> bool:
        return (
            self.recall >= UMBRALES["recall"]
            and self.mrr >= UMBRALES["mrr"]
            and self.negativas >= UMBRALES["negativas"]
        )

    def como_dict(self) -> dict[str, Any]:
        return {
            "modo": self.modo,
            f"recall@{K}": round(self.recall, 3),
            "mrr": round(self.mrr, 3),
            "negativas": f"{self.negativas}/{self.total_negativas}",
            "fallidas": self.fallidas,
            "negativas_con_resultado": self.negativas_con_resultado,
        }


# ------------------------------------------------------------------ documentos
def _secciones_md(texto: str, nivel: str) -> list[tuple[str, str]]:
    salida: list[tuple[str, str]] = []
    for bloque in re.split(rf"^{re.escape(nivel)} ", texto, flags=re.M)[1:]:
        titulo, _, cuerpo = bloque.partition("\n")
        salida.append((titulo.strip(), cuerpo.strip()))
    return salida


def _pdf(paginas: list[str]) -> bytes:
    """PDF mínimo con una página por texto, en Helvetica WinAnsi (tildes y eñes legibles)."""
    objetos: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids: list[str] = []
    fuente = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    for texto in paginas:
        n_pagina = len(objetos) + 1
        kids.append(f"{n_pagina} 0 R")
        linea = texto.replace("\\", "\\\\").replace("(", r"\(").replace(")", r"\)").encode("cp1252")
        stream = b"BT /F1 10 Tf 40 760 Td (" + linea + b") Tj ET"
        objetos.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n_pagina + 1} 0 R "
            "/Resources << /Font << /F1 ".encode()
            + fuente
            + b" >> >> >>"
        )
        objetos.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objetos[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(paginas)} >>".encode()
    salida = io.BytesIO()
    salida.write(b"%PDF-1.4\n")
    offsets = []
    for i, cuerpo in enumerate(objetos, start=1):
        offsets.append(salida.tell())
        salida.write(f"{i} 0 obj\n".encode() + cuerpo + b"\nendobj\n")
    xref = salida.tell()
    salida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        salida.write(f"{off:010d} 00000 n \n".encode())
    salida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return salida.getvalue()


def _docx(secciones: list[tuple[str, str]]) -> bytes:
    from docx import Document

    documento = Document()
    for titulo, cuerpo in secciones:
        documento.add_heading(titulo, level=1)
        for parrafo in cuerpo.split("\n\n"):
            documento.add_paragraph(parrafo.strip())
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


def documentos() -> list[tuple[str, str, bytes, str]]:
    """`(título, mime, contenido, nombre)` de los tres documentos del golden set."""
    from pack_knowledge.archivos import MIME_DOCX, MIME_PDF, MIME_TXT

    fuentes = DIR / "documentos"
    faq = _secciones_md((fuentes / "preguntas-frecuentes.md").read_text(encoding="utf-8"), "##")
    politica = _secciones_md((fuentes / "politica-de-cambios.md").read_text(encoding="utf-8"), "#")
    return [
        ("Preguntas frecuentes", MIME_PDF, _pdf([f"{t}. {c}" for t, c in faq]), "preguntas-frecuentes.pdf"),
        ("Política de cambios", MIME_DOCX, _docx(politica), "politica-de-cambios.docx"),
        ("Catálogo", MIME_TXT, (fuentes / "catalogo.txt").read_bytes(), "catalogo.txt"),
    ]


def preguntas() -> list[dict[str, str]]:
    return json.loads((DIR / "preguntas.json").read_text(encoding="utf-8"))["preguntas"]


# ------------------------------------------------------------------ ingesta y medición
@contextmanager
def _cola_en_linea() -> Iterator[list[tuple[str, dict[str, Any]]]]:
    """Las tareas del worker se capturan y se corren en orden (sin broker)."""
    from pack_knowledge import ingesta

    encoladas: list[tuple[str, dict[str, Any]]] = []
    original = ingesta.task_queue.enqueue
    ingesta.task_queue.enqueue = lambda nombre, **kw: encoladas.append((nombre, kw))  # type: ignore[method-assign]
    try:
        yield encoladas
    finally:
        ingesta.task_queue.enqueue = original  # type: ignore[method-assign]


async def ingerir(tenant_id: uuid.UUID, agent_id: uuid.UUID) -> None:
    """Sube y procesa los tres documentos en el negocio `tenant_id` (dentro de su scope)."""
    from pack_knowledge import ingesta

    from core.shared.tenant import tenant_scope

    tareas = {ingesta.TAREA_EXTRAER: ingesta.tarea_extraer, ingesta.TAREA_EMBEBER: ingesta.tarea_embeber}
    with _cola_en_linea() as encoladas:
        async with tenant_scope(tenant_id):
            for titulo, mime, contenido, _ in documentos():
                await ingesta.crear_documento(
                    agent_id=agent_id, titulo=titulo, tipo=mime, contenido=contenido, created_by=None
                )
        while encoladas:
            nombre, kwargs = encoladas.pop(0)
            await tareas[nombre](**kwargs)
    async with tenant_scope(tenant_id):
        estados = {d.titulo: d.estado for d in await ingesta.listar(agent_id)}
    if set(estados.values()) != {"listo"}:
        raise RuntimeError(f"documentos sin procesar: {estados}")


def _normal(texto: str) -> str:
    plano = unicodedata.normalize("NFKD", texto.casefold())
    return "".join(c for c in plano if not unicodedata.combining(c))


async def medir(tenant_id: uuid.UUID, agent_id: uuid.UUID, *, modo: str = "hibrida") -> Metricas:
    from pack_knowledge.busqueda import Config, buscar

    from core.shared.tenant import tenant_scope

    config = Config(modo=modo, top_k=K)  # type: ignore[arg-type]
    aciertos = 0
    rr = 0.0
    fallidas: list[str] = []
    negativas = 0
    total_negativas = 0
    negativas_con_resultado: list[str] = []
    async with tenant_scope(tenant_id):
        for p in preguntas():
            resultado = await buscar(agent_id, p["pregunta"], config)
            if "clave" not in p:
                total_negativas += 1
                if resultado.fragmentos:
                    negativas_con_resultado.append(p["id"])
                else:
                    negativas += 1
                continue
            rango = next(
                (
                    i
                    for i, f in enumerate(resultado.fragmentos, start=1)
                    if f.documento == p["documento"] and _normal(p["clave"]) in _normal(f.texto)
                ),
                None,
            )
            if rango is None:
                fallidas.append(p["id"])
            else:
                aciertos += 1
                rr += 1 / rango
    con_respuesta = len(preguntas()) - total_negativas
    return Metricas(
        modo=modo,
        recall=aciertos / con_respuesta,
        mrr=rr / con_respuesta,
        negativas=negativas,
        total_negativas=total_negativas,
        fallidas=fallidas,
        negativas_con_resultado=negativas_con_resultado,
    )


# ------------------------------------------------------------------ CLI
async def _negocio_de_prueba() -> tuple[uuid.UUID, uuid.UUID]:
    from sqlalchemy import text

    from core.alta_negocio import onboard_tenant
    from core.db.session import get_session
    from core.shared.tenant import tenant_scope

    slug = f"eval-conocimiento-{uuid.uuid4().hex[:8]}"
    tenant_id, _ = await onboard_tenant(name="Café La Montaña (eval)", slug=slug, persona={"tone": "cercano"})
    agent_id = uuid.uuid4()
    async with tenant_scope(tenant_id):
        async with get_session() as session:
            await session.execute(
                text("insert into agents (id, tenant_id, key, kind, name) values (:id, :t, 'ventas', 'specialist', 'Ventas')"),
                {"id": agent_id, "t": tenant_id},
            )
    return tenant_id, agent_id


async def _correr(args: argparse.Namespace) -> int:
    from core.llm.embeddings import grabacion

    tenant_id, agent_id = await _negocio_de_prueba()
    if args.hash:
        contexto: Any = _sin_grabacion()
    else:
        modo = "record" if args.grabar else "replay"
        contexto = grabacion(GRABACION, modo, api_key=os.environ.get("OPENAI_API_KEY") if args.grabar else None)
    with contexto:
        await ingerir(tenant_id, agent_id)
        resultados = [await medir(tenant_id, agent_id, modo=m) for m in ("hibrida", "semantica")]
    print(json.dumps([r.como_dict() for r in resultados], ensure_ascii=False, indent=2))
    if args.hash:
        return 0
    hibrida, semantica = resultados
    ok = hibrida.cumple() and hibrida.recall >= semantica.recall
    print("OK" if ok else f"NO CUMPLE: umbrales {UMBRALES} y la híbrida no por debajo de la semántica")
    return 0 if ok else 1


@contextmanager
def _sin_grabacion() -> Iterator[None]:
    yield


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    grupo = parser.add_mutually_exclusive_group()
    grupo.add_argument("--grabar", action="store_true", help="llamada real a OpenAI (con OK del dueño)")
    grupo.add_argument("--hash", action="store_true", help="embeddings hash, sin grabación")
    args = parser.parse_args()
    if args.grabar and not os.environ.get("OPENAI_API_KEY"):
        print("Falta OPENAI_API_KEY para grabar.", file=sys.stderr)
        return 2
    return asyncio.run(_correr(args))


if __name__ == "__main__":
    sys.exit(main())
